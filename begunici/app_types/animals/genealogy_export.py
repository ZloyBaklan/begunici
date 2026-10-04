from io import BytesIO
import re
from urllib.parse import quote

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import serializers
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .models import Ewe, Maker, Ram, Sheep


ANIMAL_TYPES = (
    (Maker, "maker", "Баран-производитель"),
    (Ram, "ram", "Баранчик"),
    (Ewe, "ewe", "Ярка"),
    (Sheep, "sheep", "Овцематка"),
)
ANCESTOR_COLUMNS = (
    "О", "ОМ", "ОО", "ООО", "ООМ", "ОМО", "ОММ",
    "ОООО", "ОООМ", "ООМО", "ООММ", "ОМОО", "ОМОМ", "ОММО", "ОМММ",
    "М", "ММ", "МО", "МММ", "ММО", "МОМ", "МОО",
    "ММММ", "МММО", "ММОМ", "МОМО", "ММОО", "МОММ", "МООМ", "МООО",
)
GENEALOGY_HEADERS = (
    "Половозрастная группа", "Кличка", "Номер животного", "Электронная метка",
    "Дата рождения", "Линия", *ANCESTOR_COLUMNS,
)
PAGE_SIZE = 50


def load_animal_catalog(*, include_archived):
    """Load parents in four queries, rather than querying each pedigree cell."""
    catalog = []
    fields = (
        "tag_id", "tag__tag_number", "birth_date", "mother", "father",
        "rshn_tag", "is_archived", "animal_status__status_type",
    )
    for model, type_code, label in ANIMAL_TYPES:
        animals = model.objects.all()
        if not include_archived:
            animals = animals.filter(is_archived=False)
        model_fields = (*fields, "name") if model is Maker else fields
        for animal in animals.values(*model_fields).order_by("tag__tag_number", "tag_id"):
            animal["animal_type"] = type_code
            animal["animal_type_label"] = label
            catalog.append(animal)
    return catalog


class GenealogySearchSerializer(serializers.Serializer):
    search = serializers.CharField(required=False, allow_blank=True, max_length=2000)
    page = serializers.IntegerField(required=False, default=1, min_value=1)


class GenealogySelectionSerializer(serializers.Serializer):
    tag_ids = serializers.ListField(
        child=serializers.IntegerField(min_value=1), allow_empty=False, max_length=10000,
        error_messages={"empty": "Выберите хотя бы одно животное"},
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def genealogy_animals_api(request):
    serializer = GenealogySearchSerializer(data=request.query_params)
    serializer.is_valid(raise_exception=True)
    terms = [
        term.strip().casefold()
        for term in re.split(r"[,;\n]+", serializer.validated_data.get("search", ""))
        if term.strip()
    ]
    animals = load_animal_catalog(include_archived=False)
    if terms:
        animals = [animal for animal in animals if any(
            term in str(animal.get(field) or "").casefold()
            for term in terms for field in ("tag__tag_number", "rshn_tag", "name")
        )]
    count = len(animals)
    num_pages = max(1, (count + PAGE_SIZE - 1) // PAGE_SIZE)
    page = min(serializer.validated_data["page"], num_pages)
    results = []
    for animal in animals[(page - 1) * PAGE_SIZE:page * PAGE_SIZE]:
        tag = animal["tag__tag_number"]
        name = animal.get("name")
        results.append({
            "tag_id": animal["tag_id"],
            "tag_number": tag,
            "animal_type": animal["animal_type"],
            "animal_type_label": animal["animal_type_label"],
            "display_name": f"{name}({tag})" if name else tag,
            "status": animal["animal_status__status_type"] or "Не определено",
        })
    return Response({"results": results, "count": count, "page": page, "num_pages": num_pages})


def build_parent_lookup(catalog):
    exact = {animal["tag__tag_number"]: animal for animal in catalog}
    folded = {}
    for tag, animal in exact.items():
        key = tag.casefold()
        # An ambiguous legacy spelling must not silently resolve to the wrong animal.
        folded[key] = animal if key not in folded else None
    return exact, folded


def build_genealogy_row(animal, exact, folded):
    ancestors = []
    for path in ANCESTOR_COLUMNS:
        current = animal
        tag = None
        # Paths are bounded to four generations, including for cyclic legacy data.
        for relation in path:
            if current is None:
                tag = None
                break
            tag = str(current.get("father" if relation == "О" else "mother") or "").strip()
            current = exact.get(tag) or folded.get(tag.casefold())
        ancestors.append(tag or "-")
    return [
        animal["animal_type_label"], animal.get("name") or "-",
        animal["tag__tag_number"], animal["rshn_tag"] or "-",
        animal["birth_date"] or "-", "-", *ancestors,
    ]


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def genealogy_export_excel(request):
    serializer = GenealogySelectionSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    tag_ids = list(dict.fromkeys(serializer.validated_data["tag_ids"]))
    catalog = load_animal_catalog(include_archived=True)
    active = {animal["tag_id"]: animal for animal in catalog if not animal["is_archived"]}
    if any(tag_id not in active for tag_id in tag_ids):
        return Response({
            "error": "Некоторые выбранные животные удалены или перенесены в архив. Обновите выбор"
        }, status=400)
    exact, folded = build_parent_lookup(catalog)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Генеалогия"
    sheet.append(GENEALOGY_HEADERS)
    for tag_id in tag_ids:
        sheet.append(build_genealogy_row(active[tag_id], exact, folded))
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row > 1 and cell.column == 5 and cell.value != "-":
                cell.number_format = "dd.mm.yyyy"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E9ECEF")
    sheet.row_dimensions[1].height = 32
    widths = (24, 22, 20, 24, 17, 14)
    for column in range(1, len(GENEALOGY_HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = widths[column - 1] if column <= 6 else 17
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    filename = "genealogy.xlsx"
    if len(tag_ids) == 1:
        tag_number = active[tag_ids[0]]["tag__tag_number"]
        safe_tag = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", tag_number).strip(" .") or "animal"
        filename = f"genealogy_{safe_tag}.xlsx"
    fallback_name = filename if filename.isascii() else "genealogy.xlsx"
    response["Content-Disposition"] = f"attachment; filename=\"{fallback_name}\"; filename*=UTF-8''{quote(filename)}"
    return response
