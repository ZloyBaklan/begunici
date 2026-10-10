"""Read-only animal bridge and one Decimal calculation for UI, stock and Excel."""
from calendar import monthrange
from copy import copy
from datetime import date
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from django.http import HttpResponse
from django.utils import timezone

from begunici.app_types.animals import feed_plan as legacy
from .models import FeedNorm, FeedProduct, FeedSettings


def number(value):
    result = format(value, "f")
    return result.rstrip("0").rstrip(".") if "." in result else result


def build_plan(as_of_date=None):
    as_of_date = as_of_date or timezone.localdate()
    # Read revision before the rows: concurrent edits can only make this form
    # stale, never give older norm values a newer version accepted on save.
    revision = FeedSettings.objects.get(pk=1).revision
    counts = dict.fromkeys(legacy.FEED_GROUP_LABELS, 0)
    unclassified = 0
    for model in (legacy.Maker, legacy.Sheep, legacy.Ewe, legacy.Ram):
        for animal in legacy.get_active_feed_queryset(model):
            group = legacy.classify_feed_plan_animal(animal, as_of_date)
            if group:
                counts[group] += 1
            else:
                unclassified += 1
    products = list(FeedProduct.objects.all())
    norms = {(n.product_id, n.group): n.per_head for n in FeedNorm.objects.all()}
    days = monthrange(as_of_date.year, as_of_date.month)[1]
    year_end = date(as_of_date.year, 12, 31)
    year_days = (year_end - as_of_date).days + 1
    rows, totals = [], {p.pk: Decimal(0) for p in products}
    missing = 0
    for group, label in legacy.FEED_GROUP_LABELS.items():
        cells = []
        for product in products:
            norm = norms.get((product.pk, group))
            daily = (norm or Decimal(0)) * counts[group]
            totals[product.pk] += daily
            missing += int(norm is None and counts[group] > 0)
            cells.append({"product_id": product.pk, "product_name": product.name,
                          "norm": "" if norm is None else number(norm),
                          "daily": number(daily), "monthly": number(daily * days),
                          "yearly": number(daily * year_days)})
        rows.append({"group": group, "label": label, "count": counts[group], "cells": cells})
    return {
        "date": as_of_date.isoformat(), "date_label": as_of_date.strftime("%d.%m.%Y"),
        "month": legacy.MONTH_NAMES[as_of_date.month],
        "year": as_of_date.year, "days": days, "rows": rows,
        "year_end_label": year_end.strftime("%d.%m.%Y"), "year_days": year_days,
        "heads": sum(counts.values()), "unclassified": unclassified, "missing": missing,
        "revision": revision,
        "products": [{"id": p.pk, "name": p.name, "unit": p.get_unit_display(),
                      "package": p.package, "daily": number(totals[p.pk]),
                      "monthly": number(totals[p.pk] * days),
                      "yearly": number(totals[p.pk] * year_days)} for p in products],
    }


def generate_workbook(plan=None, period="monthly"):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    plan = plan or build_plan()
    if period not in {"monthly", "yearly"}:
        raise ValueError("Unknown feed plan period")
    yearly = period == "yearly"
    days = plan["year_days"] if yearly else plan["days"]
    workbook = load_workbook(legacy.get_template_path())
    sheet = workbook.active
    sheet.title = f"до конца {plan['year']} года" if yearly else f"кормовой план {plan['month']}"
    sheet["B3"] = "до конца года" if yearly else f"на {plan['month']}"
    sheet["E3"] = f"{plan['year']} г."
    sheet["A20"] = (f"Период: {plan['date_label']}–{plan['year_end_label']} включительно, {days} дн."
                    if yearly else f"По текущему поголовью на {plan['date']}. Прогноз на {days} дней.")
    sheet["A21"] = f"Без категории: {plan['unclassified']}. Незаполненных норм при наличии голов: {plan['missing']}. Пустые нормы в итог не входят."
    if yearly:
        sheet["A22"] = "Прогноз по текущему поголовью и нормам; будущие изменения состава стада не учитываются."
    cache = {}
    for i, product in enumerate(plan["products"]):
        col = 5 + 2 * i
        left, right = get_column_letter(col), get_column_letter(col + 1)
        # Extend the established two-column layout for user-created feed items.
        if col > 17:
            for row in range(6, 19):
                for offset in (0, 1):
                    sheet.cell(row, col + offset)._style = copy(sheet.cell(row, 17 + offset)._style)
            sheet.merge_cells(start_row=6, start_column=col, end_row=7, end_column=col + 1)
            sheet.column_dimensions[left].width = sheet.column_dimensions["Q"].width
            sheet.column_dimensions[right].width = sheet.column_dimensions["R"].width
        sheet.cell(6, col, product["name"])
        sheet.cell(6, col).data_type = "s"  # User-entered names are never Excel formulas.
        sheet.cell(8, col, "на 1 голову / день")
        sheet.cell(8, col + 1, "на поголовье / до конца года" if yearly else "на поголовье / месяц")
        sheet.cell(9, col, product["unit"])
        sheet.cell(9, col + 1, product["unit"])
        sheet.column_dimensions[left].width = max(sheet.column_dimensions[left].width or 0, 11)
        sheet.column_dimensions[right].width = max(sheet.column_dimensions[right].width or 0, 14)
        for row, category in enumerate(plan["rows"], 10):
            cell = category["cells"][i]
            sheet[f"{left}{row}"] = Decimal(cell["norm"]) if cell["norm"] else None
            sheet[f"{right}{row}"] = f"=B{row}*{left}{row}*C{row}"
            sheet[f"{left}{row}"].number_format = "#,##0.000"
            sheet[f"{right}{row}"].number_format = "#,##0.000"
            cache[f"{right}{row}"] = cell[period]
        sheet[f"{right}18"] = f"=SUM({right}10:{right}17)"
        sheet[f"{right}18"].number_format = "#,##0.000"
        cache[f"{right}18"] = product[period]
    for row, category in enumerate(plan["rows"], 10):
        sheet[f"B{row}"], sheet[f"C{row}"] = category["count"], days
        sheet[f"D{row}"] = f"=B{row}*C{row}"
        cache[f"D{row}"] = str(category["count"] * days)
    sheet["B18"], sheet["D18"] = "=SUM(B10:B17)", "=SUM(D10:D17)"
    cache["B18"], cache["D18"] = str(plan["heads"]), str(plan["heads"] * days)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    sheet.print_area = f"A3:{get_column_letter(4 + 2 * len(plan['products']))}{22 if yearly else 21}"
    stream = BytesIO()
    workbook.save(stream)
    # openpyxl deliberately does not calculate. Supply values from the same plan
    # while retaining Excel formulas so previews and subsequent edits both work.
    output = BytesIO()
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(stream) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                for cell in root.iter(ns + "c"):
                    if cell.get("r") in cache and cell.find(ns + "f") is not None:
                        value = cell.find(ns + "v")
                        if value is None:
                            value = ET.SubElement(cell, ns + "v")
                        value.text = cache[cell.get("r")]
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
    output.seek(0)
    return output


def feed_plan_response(as_of_date=None, period="monthly"):
    plan = build_plan(as_of_date)
    response = HttpResponse(generate_workbook(plan, period=period).getvalue(), content_type=
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    suffix = "_do_kontsa_goda" if period == "yearly" else ""
    response["Content-Disposition"] = f'attachment; filename="kormovoy_plan_{plan["date"]}{suffix}.xlsx"'
    return response
