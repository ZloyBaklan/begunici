"""Export adapters use the project's existing openpyxl/template convention.

The source workbooks are only read. Every generated version is a separate byte
stream and is frozen in DocumentVersion before the user can approve it.
"""
from collections import defaultdict
from copy import copy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from textwrap import wrap

from django.conf import settings
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import ColBreak, RowBreak
from openpyxl.utils import get_column_letter

from .models import Product


def text_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"  # User-entered strings never become Excel formulas.
    return cell


def table_sheet(workbook, name, headers, rows):
    sheet = workbook.create_sheet(name)
    for column, title in enumerate(headers, 1):
        cell = text_cell(sheet, 1, column, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="285A47")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = 23 if column < 5 else 32
    sheet.row_dimensions[1].height = 32
    for row_index, values in enumerate(rows, 2):
        for column, value in enumerate(values, 1):
            cell = text_cell(sheet, row_index, column, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(value, (float, Decimal)):
                cell.number_format = '#,##0.000'
        sheet.row_dimensions[row_index].height = 34
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.print_options.horizontalCentered = True
    return sheet


def numeric(value):
    return Decimal(str(value)) if value not in (None, "") else None


def generate_document(document):
    if document.kind == "sp55":
        workbook = slaughter_workbook(document)
    elif document.kind == "sp54":
        workbook = live_workbook(document)
    else:
        workbook = invoice_workbook(document)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_slaughter_output_rows(sheet, meat_rows, offal_rows, meat_weight, offal_weight):
    rows = [(name, weight, False) for name, weight in meat_rows]
    if not rows:
        rows.append(("Мясо", Decimal(0), False))
    rows.extend([
        ("Итого мясо", meat_weight, True),
        ("Субпродукты по видам:", None, False),
    ])
    for name, items in offal_rows.items():
        rows.extend((f"{name} ({tag})", weight, False) for tag, weight in items)
        rows.append((f"Итого {name.lower()}", sum((weight for _, weight in items), Decimal(0)), True))
    rows.append(("Итого субпродукты", offal_weight, True))

    styles = {column: copy(sheet.cell(19, column)._style) for column in range(17, 33)}
    extra_rows = max(0, len(rows) - 9)
    # Inserting cells alone does not move merged ranges or row dimensions.
    footer_merges = [copy(merged) for merged in sheet.merged_cells.ranges if merged.min_row >= 25]
    body_merges = [str(merged) for merged in sheet.merged_cells.ranges
                   if merged.min_row >= 16 and merged.max_row <= 24 and merged.min_col >= 17]
    for merged in body_merges:
        sheet.unmerge_cells(merged)
    if extra_rows:
        footer_dimensions = [(index, copy(dimension)) for index, dimension in sheet.row_dimensions.items() if index >= 25]
        for merged in footer_merges:
            sheet.unmerge_cells(str(merged))
        for index, _ in footer_dimensions:
            del sheet.row_dimensions[index]
        sheet.insert_rows(25, amount=extra_rows)
        for index, dimension in footer_dimensions:
            dimension.index = index + extra_rows
            sheet.row_dimensions[dimension.index] = dimension
        for merged in footer_merges:
            merged.shift(row_shift=extra_rows)
            sheet.merge_cells(str(merged))

    for row_index in range(16, 25 + extra_rows):
        for column, style in styles.items():
            cell = sheet.cell(row_index, column)
            cell.value = None
            cell._style = copy(style)
        for first_column, last_column in ((18, 21), (22, 23), (25, 28), (29, 32)):
            sheet.merge_cells(start_row=row_index, start_column=first_column,
                              end_row=row_index, end_column=last_column)
        sheet.row_dimensions[row_index].height = max(sheet.row_dimensions[row_index].height or 0, 30)
    for row_index, (name, weight, is_total) in enumerate(rows, 16):
        label = text_cell(sheet, row_index, 18, name)
        label.alignment = Alignment(wrap_text=True, vertical="center")
        font = copy(label.font)
        font.bold = is_total
        label.font = font
        sheet.cell(row_index, 24, weight)
        # Leave enough height for long product names and identifiers, without shrinking text.
        sheet.row_dimensions[row_index].height = max(30, ((len(name) + 17) // 18) * 12 + 6)
    total_row = 25 + extra_rows
    total_label = text_cell(sheet, total_row, 18, "Итого мясо и субпродукты")
    total_label.alignment = Alignment(wrap_text=True, vertical="center")
    total_font = copy(total_label.font)
    total_font.bold = True
    total_label.font = total_font
    sheet.row_dimensions[total_row].height = max(sheet.row_dimensions[total_row].height or 0, 42)
    sheet.cell(total_row, 24, meat_weight + offal_weight)
    for row_index in [*range(16, total_row), total_row]:
        sheet.cell(row_index, 24).number_format = '#,##0.000'
        sheet.cell(row_index, 24).alignment = Alignment(horizontal="center", vertical="center")
    sheet.print_area = f"A1:AF{33 + extra_rows}"
    sheet.print_title_rows = "1:15"
    sheet.page_setup.fitToHeight = 0


def _sp55_print_state(sheet):
    return (str(sheet.print_area), copy(sheet.page_setup), copy(sheet.page_margins),
            copy(sheet.print_options), copy(sheet.sheet_properties.pageSetUpPr),
            sheet.sheet_view.showGridLines, str(sheet.oddFooter), str(sheet.oddHeader),
            copy(sheet.row_breaks), copy(sheet.col_breaks))


def _sp55_print_settings(sheet, area, fit_height=0):
    before = _sp55_print_state(sheet)
    sheet.print_area = area
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.scale = None
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = fit_height
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.25, bottom=0.3,
                                     header=0.1, footer=0.15)
    sheet.print_options.horizontalCentered = True
    sheet.print_options.gridLines = False
    sheet.sheet_view.showGridLines = False
    sheet.oddFooter.right.text = "Стр. &P"
    sheet.oddFooter.right.size = 8
    sheet.oddHeader.right.text = None
    sheet.row_breaks = RowBreak()
    sheet.col_breaks = ColBreak()
    return before != _sp55_print_state(sheet)


def _sp55_merge(sheet, area, value):
    from openpyxl.worksheet.cell_range import CellRange

    target = CellRange(area)
    for merged in list(sheet.merged_cells.ranges):
        if not merged.isdisjoint(target):
            sheet.unmerge_cells(str(merged))
    for row in sheet.iter_rows(min_row=target.min_row, max_row=target.max_row,
                               min_col=target.min_col, max_col=target.max_col):
        for cell in row:
            cell.value = None
    sheet.merge_cells(area)
    text_cell(sheet, target.min_row, target.min_col, value)


def _sp55_row_height(sheet, row, fields, minimum=24):
    height = minimum
    for first, last, value in fields:
        # Excel cannot auto-fit merged cells. Reserve height for explicit and
        # wrapped lines, including long unbroken tags, without shrinking fonts.
        width = sum(sheet.column_dimensions[get_column_letter(col)].width
                    for col in range(first, last + 1))
        capacity = max(1, int(width * 0.88))
        lines = sum(max(1, len(wrap(part, width=capacity)))
                    for part in str(value or "").split("\n"))
        height = max(height, lines * 13 + 5)
    sheet.row_dimensions[row].height = min(409, height)


def _sp55_template_header(sheet, reference):
    changed = False
    subtitle = str(reference["L5"].value or "").strip()
    if "K5:W6" not in sheet.merged_cells or sheet["K5"].value != subtitle:
        _sp55_merge(sheet, "K5:W6", subtitle)
        changed = True
    for address, source_address in (("K3", "K3"), ("K5", "L5"), ("D7", "D7")):
        font = copy(reference[source_address].font)
        if address in {"K3", "K5"}:
            font.bold = True
            font.italic = False
        if sheet[address].font != font:
            sheet[address].font = font
            changed = True
        alignment = (Alignment(horizontal="left", vertical="center", indent=1,
                               wrap_text=False, shrink_to_fit=False)
                     if address in {"K3", "K5"} else copy(reference[source_address].alignment))
        if sheet[address].alignment != alignment:
            sheet[address].alignment = alignment
            changed = True
    if sheet["D7"].value != reference["D7"].value:
        text_cell(sheet, 7, 4, reference["D7"].value)
        changed = True
    value = sheet["AD6"].value
    if isinstance(value, str) and "." in value:
        for pattern in ("%d.%m.%Y", "%d.%m.%y"):
            try:
                value = datetime.strptime(value, pattern).date()
                break
            except ValueError:
                continue
        else:
            return changed
    if isinstance(value, date):
        values = [value.strftime(part) for part in ("%d", "%m", "%y")]
    elif value in (None, "") or (sheet["AE6"].value and sheet["AF6"].value):
        values = [sheet[address].value for address in ("AD6", "AE6", "AF6")]
        year = str(values[2] or "")
        if len(year) == 4 and year.isdigit():
            values[2] = year[-2:]
    else:
        return changed
    if "AD6:AF6" in sheet.merged_cells:
        sheet.unmerge_cells("AD6:AF6")
        changed = True
    for address, value in zip(("AD6", "AE6", "AF6"), values):
        cell, source = sheet[address], reference[address]
        if cell.value != value:
            text_cell(sheet, cell.row, cell.column, value)
            changed = True
        for attribute in ("font", "border", "alignment", "number_format"):
            style = copy(getattr(source, attribute))
            if getattr(cell, attribute) != style:
                setattr(cell, attribute, style)
                changed = True
    return changed


def _sp55_finish_print(workbook, reference_workbook, total_row):
    sheet = workbook.worksheets[0]
    height = sum(sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight or 15
                 for row in range(1, total_row + 9))
    # A short act must not send just its final date/signature to the next page.
    # Long acts retain vertical pagination instead of shrinking to tiny text.
    compact = total_row <= 32 and height <= 600
    changed = _sp55_print_settings(sheet, f"A1:AF{total_row + 8}", fit_height=int(compact))
    if len(workbook.worksheets) > 1:
        reverse, reference = workbook.worksheets[1], reference_workbook.worksheets[1]
        changed |= _sp55_print_settings(reverse, "A1:X22", fit_height=1)
        if reverse.print_title_rows is not None:
            # openpyxl's public setter ignores None rather than clearing it.
            reverse._print_rows = None
            changed = True
        for key, dimension in list(reverse.column_dimensions.items()):
            if dimension.min and dimension.min > 24:
                del reverse.column_dimensions[key]
                changed = True
        for row in range(1, 22):
            if reverse.row_dimensions[row].height != reference.row_dimensions[row].height:
                reverse.row_dimensions[row].height = reference.row_dimensions[row].height
                changed = True
            for column in range(1, 25):
                cell, source = reverse.cell(row, column), reference.cell(row, column)
                for attribute in ("font", "alignment"):
                    style = copy(getattr(source, attribute))
                    if getattr(cell, attribute) != style:
                        setattr(cell, attribute, style)
                        changed = True
    for appendix in workbook.worksheets[2:]:
        if appendix.title in {"Животные", "Выход продукции", "Брак — не поступает"}:
            changed |= _sp55_print_settings(appendix, f"A1:E{appendix.max_row}")
    return changed


def format_slaughter_workbook(workbook):
    """Print projection of our SP-55 form; never recalculate frozen quantities."""
    sheet = workbook.worksheets[0]
    if (sheet.title != "СП-55(1)"
            or not str(sheet["K3"].value or "").startswith("УЧЕТНЫЙ ЛИСТ №")):
        return False
    total_row = next((row for row in range(25, sheet.max_row + 1)
                      if str(sheet.cell(row, 7).value or "").strip() == "Итого"), None)
    if total_row is None:
        return False
    reference_workbook = load_workbook(Path(settings.BASE_DIR) /
                                       "begunici/app_types/animals/excel_templates/sp-55.xlsx")
    reference = reference_workbook.worksheets[0]
    if "K3:W4" in sheet.merged_cells:
        changed = _sp55_template_header(sheet, reference)
        changed |= _sp55_finish_print(workbook, reference_workbook, total_row)
        return changed
    # AD:AF is one column-dimension group in the source. Widening AD also
    # widened AE and AF, distorting the whole form and reducing print scale.
    widths = [reference.sheet_format.defaultColWidth or 8.43] * 32
    for dimension in reference.column_dimensions.values():
        for column in range(dimension.min, min(dimension.max, 32) + 1):
            widths[column - 1] = dimension.width
    for key in list(sheet.column_dimensions):
        del sheet.column_dimensions[key]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    title = sheet["K3"].value
    old_conclusion = sheet["J15"].value
    _sp55_merge(sheet, "K3:W4", title)
    # The legacy template merged the first conclusion into its header row.
    # Read all values before removing those merges.
    body = [(row, sheet.cell(row, 5).value, sheet.cell(row, 10).value)
            for row in range(16, 25)]
    _sp55_merge(sheet, "J13:N15", sheet["J13"].value)
    for row, place, conclusion in body:
        _sp55_merge(sheet, f"A{row}:G{row}", place or "с участка № _______")
        _sp55_merge(sheet, f"J{row}:N{row}", old_conclusion if row == 16 else conclusion)
    for row in range(25, total_row):
        _sp55_merge(sheet, f"A{row}:G{row}", None)
        _sp55_merge(sheet, f"J{row}:N{row}", None)
    for column in (8, 9):
        value = sheet.cell(23, column).value
        _sp55_merge(sheet, f"{get_column_letter(column)}23:{get_column_letter(column)}24", value)
        sheet.unmerge_cells(f"{get_column_letter(column)}23:{get_column_letter(column)}24")
    thin = Side(style="thin", color="333333")
    for row in range(1, 16):
        sheet.row_dimensions[row].height = reference.row_dimensions[row].height
    _sp55_template_header(sheet, reference)
    for address in ("V12", "X12"):
        sheet[address].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[12].height = 18
    for row in range(16, total_row + 1):
        _sp55_row_height(sheet, row, [(1, 7, sheet.cell(row, 1).value),
                                    (10, 14, sheet.cell(row, 10).value),
                                    (18, 21, sheet.cell(row, 18).value)], minimum=20)
        for column in (1, 10, 18, 24, 8, 9):
            cell = sheet.cell(row, column)
            font = copy(cell.font)
            font.sz = 9.5
            cell.font = font
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal="center" if column in (8, 9, 24) else "left")
        for column in (8, 9, 24):
            sheet.cell(row, column).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row, 9).number_format = '#,##0.000'
    for row in range(16, total_row + 3):
        columns = [(17, 21), (22, 23), (24, 24), (25, 28), (29, 32)]
        if row < total_row:
            columns += [(1, 7), (8, 8), (9, 9), (10, 14)]
        for first, last in columns:
            for column in range(first, last + 1):
                sheet.cell(row, column).border = Border(
                    top=thin, bottom=thin, left=thin if column == first else Side(),
                    right=thin if column == last else Side())
    for row in (total_row, total_row + 1, total_row + 2):
        _sp55_row_height(sheet, row, [(18, 21, sheet.cell(row, 18).value)], minimum=20)
        sheet.cell(row, 7).alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
        sheet.cell(row, 9).number_format = '#,##0.000'
        for column in (8, 9, 24):
            font = copy(sheet.cell(row, column).font)
            font.sz = 9.5
            sheet.cell(row, column).font = font
    for row in range(total_row + 3, total_row + 9):
        source_row = 28 + row - (total_row + 3)
        sheet.row_dimensions[row].height = reference.row_dimensions[source_row].height
    sheet.print_title_rows = "1:15"
    for appendix in workbook.worksheets[2:]:
        if appendix.title not in {"Животные", "Выход продукции", "Брак — не поступает"}:
            continue
        _sp55_print_settings(appendix, f"A1:E{appendix.max_row}")
        for column, width in enumerate([19, 32, 14, 17, 54], 1):
            appendix.column_dimensions[get_column_letter(column)].width = width
        for row in appendix.iter_rows(max_col=5):
            for cell in row:
                font = copy(cell.font)
                font.name, font.sz, font.scheme = "Times New Roman", 10, None
                cell.font = font
            _sp55_row_height(appendix, row[0].row, [(cell.column, cell.column, cell.value) for cell in row])
        appendix.print_title_rows = "1:1"
    _sp55_finish_print(workbook, reference_workbook, total_row)
    return True


def printable_slaughter_excel(content):
    try:
        workbook = load_workbook(BytesIO(content))
    except Exception:
        return content
    if not format_slaughter_workbook(workbook):
        return content
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def slaughter_workbook(document):
    template = Path(settings.BASE_DIR) / "begunici/app_types/animals/excel_templates/sp-55.xlsx"
    workbook = load_workbook(template)
    sheet = workbook.worksheets[0]
    sheet["K3"] = f"УЧЕТНЫЙ ЛИСТ № {getattr(document, 'act_number', None) or document.number or document.pk}"
    sheet["AD6"] = document.date
    sheet["AD6"].number_format = "dd.mm.yyyy"
    sheet["E8"] = "Убойный цех"
    sheet["F9"] = "См. строки животных и приложение"
    sheet["V12"] = "Склад продукции"
    sheet["X12"] = "Убойный цех"
    sheet["J28"] = "Овец в живой массе принял"
    workbook.worksheets[1]["B14"] = "Убой овец указанной массы и выход"
    animals = list(document.animals.all())
    animal_rows = []
    for case in animals:
        details = document.payload.get("animals", {}).get(str(case.pk), {})
        conclusion = "; ".join(filter(None, [details.get("fatness", ""), details.get("comment", "")]))
        animal_rows.append([
            case.tag_number, case.snapshot.get("place", ""), 1,
            numeric(details.get("weight")), conclusion,
        ])
    # Fixed source form has 8 rows. All per-animal data remains in the attached
    # dynamic table; overflow is summarized instead of cutting records off.
    groups = defaultdict(list)
    for item in animal_rows:
        groups[item[1]].append(item)
    display = list(groups.items())
    if len(display) > 8:
        display = [("Все участки — см. приложение", animal_rows)]
    for index, (place, members) in enumerate(display):
        row = 16 + index
        text_cell(sheet, row, 5, place)
        sheet[f"H{row}"] = len(members)
        sheet[f"I{row}"] = sum((member[3] or Decimal(0) for member in members), Decimal(0))
        target = "J15" if row == 16 else f"J{row}"
        conclusion = "; ".join(f"{member[0]}: {member[4]}" for member in members)
        if len(conclusion) > 250:
            conclusion = "Бирки и заключения ветврача - см. лист «Животные»"
        text_cell(sheet, sheet[target].row, sheet[target].column, conclusion)
        sheet[target].alignment = Alignment(wrap_text=True, vertical="top")
    sheet["H25"] = len(animals)
    sheet["H26"] = len(animals)
    sheet["I25"] = sum((row[3] or Decimal(0) for row in animal_rows), Decimal(0))
    sheet["I26"] = sheet["I25"].value
    products = {p.code: p for p in Product.objects.all()}
    case_map = {case.pk: case for case in animals}
    outputs, rejected = [], []
    totals = defaultdict(Decimal)
    offal = defaultdict(list)
    meat_by_animal = defaultdict(Decimal)
    for item in document.payload.get("outputs", []):
        product = products[item["product"]]
        case = case_map[int(item["animal_id"])]
        name = item.get("name") or product.name
        category = item.get("category") or product.category
        row = [case.tag_number, name, int(item.get("quantity", 0)), numeric(item["weight"]), item.get("reason", "")]
        if item.get("rejected"):
            rejected.append(row)
        else:
            outputs.append(row)
            totals[category] += numeric(item["weight"])
            if category == "meat":
                meat_by_animal[case.pk] += numeric(item["weight"])
            elif category == "offal":
                offal[name].append((case.tag_number, numeric(item["weight"])))
    # Carcass mass is separate from live mass. Never add the entered cuts to
    # their full carcass mass a second time. Rejected meat is not accepted.
    meat_weight = Decimal(0)
    meat_rows = []
    for case in animals:
        details = document.payload.get("animals", {}).get(str(case.pk), {})
        declared = numeric(details.get("carcass_weight") or case.snapshot.get("carcass_weight"))
        all_meat = [item for item in document.payload.get("outputs", [])
                    if int(item["animal_id"]) == case.pk and (item.get("category") or products[item["product"]].category) == "meat"]
        if all_meat:
            accepted_weight = meat_by_animal[case.pk]
            meat_weight += accepted_weight
            if accepted_weight:
                meat_rows.append((f"Мясо ({case.tag_number})", accepted_weight))
        elif declared is not None and not any(int(item["animal_id"]) == case.pk for item in document.payload.get("outputs", [])):
            meat_weight += declared
            meat_rows.append((f"Мясо ({case.tag_number})", declared))
    _write_slaughter_output_rows(sheet, meat_rows, offal, meat_weight, totals["offal"])
    table_sheet(workbook, "Животные", ["Бирка", "С участка", "Голов", "Масса, кг", "Заключение ветврача"], animal_rows)
    table_sheet(workbook, "Выход продукции", ["Бирка", "Номенклатура", "Количество, шт.", "Масса, кг", "Примечание"], outputs)
    table_sheet(workbook, "Брак — не поступает", ["Бирка", "Номенклатура", "Количество, шт.", "Масса, кг", "Причина брака"], rejected)
    format_slaughter_workbook(workbook)
    return workbook


def live_workbook(document):
    """Correction worksheet; initial SP-54 always comes from the original exporter.

    This accompanies structured amendments. The user supplies the corrected
    signed act; its uploaded bytes become the next version.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    case_map = {case.pk: case for case in document.animals.all()}
    rows = [[case_map[int(line["animal_id"])].tag_number, 1, numeric(line["weight"])] for line in document.payload.get("outputs", [])]
    sheet = table_sheet(workbook, "Данные СП-54", ["Бирка", "Голов", "Живая масса, кг"], rows)
    sheet.cell(len(rows) + 4, 1, f"СП-54 № {document.act_number}. Приложите исправленный СП-54 для утверждения этих данных.")
    return workbook


def invoice_workbook(document):
    from .models import StockLot
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows = []
    for item in document.payload.get("lines", []):
        lot = StockLot.objects.select_related("product", "animal").get(pk=item["lot_id"])
        rows.append([lot.animal.tag_number, item.get("name") or lot.product.name, item["quantity"], numeric(item["weight"]), item.get("order_line_id") or ""])
    sheet = table_sheet(workbook, "Товарная накладная", ["Бирка", "Номенклатура", "Количество, шт.", "Масса, кг", "Строка заявки"], rows)
    text_cell(sheet, len(rows) + 4, 1, f"Тестовый шаблон товарной накладной № {document.number or document.pk} от {document.date:%d.%m.%Y}")
    text_cell(sheet, len(rows) + 5, 1, f"Получатель: {document.payload.get('customer', '')}")
    text_cell(sheet, len(rows) + 6, 1, f"Заявка: {document.order_id or 'не связана'}")
    sheet.merge_cells(start_row=len(rows) + 4, start_column=1, end_row=len(rows) + 4, end_column=5)
    return workbook
