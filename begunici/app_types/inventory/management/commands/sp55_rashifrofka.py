"""One-off SP-55 Excel update. Preview: python manage.py sp55_rashifrofka

Apply: python manage.py sp55_rashifrofka --apply
Only content and sha256 are replaced. Frozen snapshots and stock are untouched.
No backup files, audit events or user action logs are created.
"""

from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO

from django.core.management.base import BaseCommand, CommandError


class UnsafeVersion(ValueError):
    pass


def digest(content):
    return hashlib.sha256(content).hexdigest()


def weight(value):
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise UnsafeVersion(f"Некорректный исторический вес: {value!r}") from None
    if not result.is_finite() or result < 0:
        raise UnsafeVersion(f"Некорректный исторический вес: {value!r}")
    return result


def appendix_rows(sheet):
    expected = ("Бирка", "Номенклатура", "Количество, шт.", "Масса, кг", "Примечание")
    if tuple(cell.value for cell in sheet[1][:5]) != expected:
        # The rejection appendix has a different label in its last column.
        if sheet.title != "Брак — не поступает" or tuple(
            cell.value for cell in sheet[1][:5]
        ) != (
            *expected[:4],
            "Причина брака",
        ):
            raise UnsafeVersion(f"Неизвестные колонки листа {sheet.title}")
    result = []
    for values in sheet.iter_rows(min_row=2, max_col=5, values_only=True):
        if all(value is None for value in values):
            continue
        tag, name, quantity, value, note = values
        result.append(
            (str(tag), str(name), int(quantity), weight(value), str(note or ""))
        )
    return Counter(result)


def detailed_content(content, snapshot):
    """Patch the saved workbook using only its frozen snapshot, never live models."""
    from openpyxl import load_workbook

    from begunici.app_types.inventory.exports import _write_slaughter_output_rows

    workbook = load_workbook(BytesIO(content))
    sheet = workbook.worksheets[0]
    if sheet.title != "СП-55(1)":
        raise UnsafeVersion("Неизвестная форма СП-55")
    if any(
        sheet.cell(row, 18).value == "Итого мясо"
        for row in range(16, sheet.max_row + 1)
    ):
        return None
    if sheet["R16"].value != "Мясо" or str(sheet["G25"].value or "").strip() != "Итого":
        raise UnsafeVersion("Структура старой формы отличается от ожидаемой")
    if not isinstance(snapshot, dict) or snapshot.get("kind") != "sp55":
        raise UnsafeVersion("Нет исторических данных СП-55")
    payload = snapshot.get("payload", {})
    animals = {
        int(item["id"]): item["snapshot"] for item in snapshot.get("animals", [])
    }
    if not animals or any(not item.get("tag_number") for item in animals.values()):
        raise UnsafeVersion("В исторических данных не сохранены бирки животных")
    outputs = payload.get("outputs", [])
    expected, rejected = [], []
    offal, meat = defaultdict(list), defaultdict(Decimal)
    animals_with_outputs, animals_with_meat = set(), set()
    for item in outputs:
        case_id = int(item["animal_id"])
        if case_id not in animals:
            raise UnsafeVersion("Выход продукции относится к неизвестному животному")
        tag = str(animals[case_id]["tag_number"])
        category = item.get("category")
        name = item.get("name")
        # Initial draft versions contain only an unnormalized carcass row.
        if item["product"] == "carcass":
            category, name = category or "meat", name or "Туша"
        if category not in {"meat", "offal"} or not name:
            raise UnsafeVersion(
                "В исторических данных не сохранены название/категория продукции"
            )
        value = weight(item["weight"])
        row = (
            tag,
            name,
            int(item.get("quantity", 0)),
            value,
            str(item.get("reason") or ""),
        )
        animals_with_outputs.add(case_id)
        if category == "meat":
            animals_with_meat.add(case_id)
        if item.get("rejected"):
            rejected.append(row)
        else:
            expected.append(row)
            if category == "meat":
                meat[case_id] += value
            else:
                offal[name].append((tag, value))
    for title, rows in (
        ("Выход продукции", expected),
        ("Брак — не поступает", rejected),
    ):
        if title not in workbook.sheetnames or appendix_rows(
            workbook[title]
        ) != Counter(rows):
            raise UnsafeVersion(
                f"Лист {title} не совпадает с историческими данными версии"
            )
    meat_rows = []
    # Early empty drafts printed zero output even when an animal's archive
    # snapshot already had a carcass weight. Preserve those historical zeros.
    empty_draft = (
        not outputs
        and weight(sheet["X16"].value) == 0
        and weight(sheet["X25"].value) == 0
    )
    for case_id, animal in animals.items():
        if empty_draft:
            value = Decimal(0)
        elif case_id in animals_with_meat:
            value = meat[case_id]
        elif case_id not in animals_with_outputs:
            details = payload.get("animals", {}).get(str(case_id), {})
            declared = details.get("carcass_weight") or animal.get("carcass_weight")
            value = weight(declared) if declared not in (None, "") else Decimal(0)
        else:
            value = Decimal(0)
        if value:
            meat_rows.append((f"Мясо ({animal['tag_number']})", value))
    meat_total = sum((value for _, value in meat_rows), Decimal(0))
    offal_total = sum(
        (value for items in offal.values() for _, value in items), Decimal(0)
    )
    if (
        weight(sheet["X16"].value) != meat_total
        or weight(sheet["X25"].value) != meat_total + offal_total
    ):
        raise UnsafeVersion(
            "Итоговые веса исходного акта не совпадают с историческими данными"
        )
    _write_slaughter_output_rows(sheet, meat_rows, offal, meat_total, offal_total)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_plan(versions):
    generated = {
        (version.document_id, version.sha256)
        for version in versions
        if version.origin == "generated"
    }
    plan, messages, errors = [], [], []
    for version in versions:
        label = (
            f"СП-55 #{version.document_id}, версия {version.number} (id={version.pk})"
        )
        if not version.filename.lower().endswith(".xlsx"):
            messages.append(f"Пропуск: {label}: не Excel")
            continue
        if (
            version.origin != "generated"
            and (version.document_id, version.sha256) not in generated
        ):
            messages.append(
                f"Пропуск: {label}: вручную загруженный файл не совпадает с формой сервиса"
            )
            continue
        try:
            original = bytes(version.content)
            if digest(original) != version.sha256:
                raise UnsafeVersion("Контрольная сумма сохранённого файла не совпадает")
            updated = detailed_content(original, version.snapshot)
            if updated is None:
                messages.append(f"Пропуск: {label}: расшифровка уже есть")
                continue
            plan.append((version, updated))
            messages.append(f"Готово к обновлению: {label}")
        except Exception as exc:
            errors.append(f"Ошибка: {label}: {exc}")
    return plan, messages, errors


def run(*, apply=False, document_ids=None, emit=print):
    from django.db import connection, transaction

    from begunici.app_types.inventory.models import (
        Document,
        DocumentVersion,
    )

    query = DocumentVersion.objects.filter(document__kind="sp55")
    if document_ids:
        query = query.filter(document_id__in=document_ids)
    with transaction.atomic():
        if apply:
            doc_ids = query.values_list("document_id", flat=True).distinct()
            list(
                Document.objects.select_for_update()
                .filter(pk__in=doc_ids)
                .order_by("pk")
            )
            query = query.select_for_update()
        versions = list(query.order_by("document_id", "number"))
        plan, messages, errors = build_plan(versions)
        for message in [*messages, *errors]:
            emit(message)
        emit(
            f"Всего версий: {len(versions)}; к обновлению: {len(plan)}; ошибок: {len(errors)}"
        )
        if errors:
            raise UnsafeVersion(
                "Обновление отменено: сначала устраните перечисленные расхождения"
            )
        if not apply:
            emit("Режим проверки: база не изменена. Для применения добавьте --apply")
            return {"updated": 0, "planned": len(plan)}
        if not plan:
            emit("Обновлять нечего; база не изменена")
            return {"updated": 0, "planned": 0}
        table = connection.ops.quote_name(DocumentVersion._meta.db_table)
        # Intentional maintenance exception to append-only history, authorized by
        # the user. Keep version identity, snapshots, metadata and stock intact.
        with connection.cursor() as cursor:
            for version, content in plan:
                new_hash = digest(content)
                cursor.execute(
                    f"UPDATE {table} SET content = %s, sha256 = %s WHERE id = %s AND sha256 = %s",
                    [content, new_hash, version.pk, version.sha256],
                )
                if cursor.rowcount != 1:
                    raise UnsafeVersion(
                        f"Версия {version.pk} изменилась во время обработки"
                    )
    emit(f"Обновлено версий: {len(plan)}. Веса, суммы и складские остатки не изменены")
    return {"updated": len(plan), "planned": len(plan)}


class Command(BaseCommand):
    help = "Расшифровка старых Excel СП-55 без изменения весов и сумм"

    def add_arguments(self, parser):
        parser.add_argument(
            "--apply",
            action="store_true",
            help="Применить обновление (по умолчанию только проверка)",
        )
        parser.add_argument(
            "--document-id",
            type=int,
            action="append",
            help="Обработать только этот акт; можно повторять",
        )

    def handle(self, *args, **options):
        try:
            run(
                apply=options["apply"],
                document_ids=options["document_id"],
                emit=self.stdout.write,
            )
        except Exception as exc:
            raise CommandError(f"{exc}. Изменения базы отменены.") from exc
